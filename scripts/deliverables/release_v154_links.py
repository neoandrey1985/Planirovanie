# -*- coding: utf-8 -*-
"""v1.54.0: стрелки зависимостей на «Ганте по иерархии» в документах и презентации. Руководство пользователя — раздел
«6. Зависимости на «Ганте по иерархии»» после «6. Метки этапов…», уточнение поля «Зависит от» (17.10); описания — раздел
«Зависимости на Ганте»; слайд после «Метки этапов»; версия во всех .docx. Снимки — capture_special_shots.py
(s_links, s_links2, s_linkpanel). Идемпотентно. Запуск: python scripts/deliverables/release_v154_links.py"""
import copy, glob, os, pathlib, docx
from docx.shared import Inches
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.53.0", "1.54.0"
STREAM_NOTE = None
H = "Зависимости на Ганте"
UG = ("6. Зависимости на «Ганте по иерархии»", "Стрелки показывают связи «окончание → начало»: от конца плановой полосы предшественника к началу полосы преемника — между задачами, историями, фичами и эпиками.",
      ["Поле «Зависит от» принимает несколько предшественников (в «Бэклоге», на «Ганте» и в панели элемента): «＋» добавляет, «×» убирает. При добавлении проверяется, что не получится цикл и что элементы не в одной ветке иерархии (родитель и так включает потомка).",
       "Цвет стрелки: серая — связь в порядке, красная — конфликт (преемник по плану заканчивается раньше предшественника), акцентная толстая — связь на критическом пути. При наведении подсвечиваются обе строки, подсказка — «T-14 ждёт T-13».",
       "Свёрнутые уровни: если элемент скрыт, связь переносится на ближайшую видимую строку — пунктирная стрелка между фичами или эпиками со значком ×N (сколько связей внутри). При раскрытии стрелки переходят на сами задачи.",
       "Связи внутри одной ветки (задача ждёт своего родителя или потомка) не рисуются — над диаграммой предупреждение, чтобы их убрать.",
       "Клик по стрелке — список связей с кнопкой «Удалить связь» (редактор). Панель элемента (клик по названию) — «Зависит от» и метки этапов.",
       "◆ после полосы — внешняя зависимость от другого стрима (раздел «Зависимости»), цвет — статус. Флажок «зависимости» над диаграммой скрывает стрелки и маркеры; все обозначения — в легенде.",
       "Критический путь, автонабор спринта, проверки данных и баннер конфликтов учитывают все предшественники."],
      [("s_links.jpg", "Стрелки зависимостей между задачами и историями"), ("s_links2.jpg", "Свёрнутый уровень: связи между фичами (×N)"), ("s_linkpanel.jpg", "Панель элемента: «Зависит от» и метки этапов")])
U1710 = ("Поля «Спринт» и «Зависит от» (Гант) — быстрое изменение размещения и связей. Как работает: меняют спринт задачи и её зависимость прямо в таблице под диаграммой; диаграмма и метрики пересчитываются.",
         "Поля «Спринт» и «Зависит от» (Гант) — быстрое изменение размещения и связей. Как работает: меняют спринт задачи и список её предшественников («＋» добавляет, «×» убирает) прямо в таблице под диаграммой; диаграмма и метрики пересчитываются.")
DESC = (H, "Стрелки зависимостей «окончание → начало» на «Ганте по иерархии» между задачами, историями, фичами и эпиками: красные — конфликт сроков, акцентные — "
           "критический путь; при свёрнутых уровнях связи переносятся на видимые фичи и эпики (пунктир, ×N). У задачи может быть несколько предшественников; "
           "связи правятся в «Бэклоге» и в панели элемента, внешние зависимости показаны ◆.")
SLIDE = ("🔗", H, "Связи «окончание → начало» на всех уровнях.", "s_links.jpg", [
    "•  Несколько предшественников у задачи, проверка циклов.",
    "•  Красные — конфликт сроков, акцент — критический путь.",
    "•  Свёрнуто — связи между фичами и эпиками (×N).",
    "•  Клик по стрелке — удалить связь; ◆ — внешние."])


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def next_heading(d, i):
    return next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))


def heading(d, text, lvl="Heading 2"):
    return next(k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.strip() == text)


fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
if not any(p.text == UG[0] for p in d.paragraphs):
    anchor = next_heading(d, heading(d, "6. Метки этапов на «Ганте по иерархии»"))
    put_before(anchor, d.add_paragraph(UG[0], style="Heading 2")); put_before(anchor, d.add_paragraph(UG[1]))
    for b in UG[2] + ([STREAM_NOTE] if STREAM_NOTE else []): put_before(anchor, d.add_paragraph(b, style="List Bullet"))
    for img, cap in UG[3]: add_pic(d, anchor, img, cap, 6.3)
    for p in d.paragraphs:
        if p.text == U1710[0]:
            for r in p.runs[1:]: r.text = ""
            p.runs[0].text = U1710[1]
    st = "добавлен"
else: st = "уже есть"
d.save(fn); print(fn[:58], "раздел:", st)
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not any(p.style.name.startswith("Heading") and p.text.strip() == H for p in d.paragraphs):
        anchor = next_heading(d, heading(d, "Метки этапов"))
        put_before(anchor, d.add_paragraph(H, style="Heading 2")); put_before(anchor, d.add_paragraph(DESC[1]))
        if pics: add_pic(d, anchor, "s_links.jpg", H, 6.61)
        st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "раздел:", st)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


DECK = "Презентация_приложения.pptx"; prs = Presentation(DECK)
s = next((x for x in prs.slides if title_of(x) == SLIDE[1]), None)
if s is None:
    i = next(k for k, x in enumerate(prs.slides) if title_of(x) == "Метки этапов")
    s = clone_after(prs, list(prs.slides)[i], i + 1, SLIDE[3])
icon, tsh, ssh, bsh, psh = parts(s)
set_text(icon, SLIDE[0]); set_text(tsh, SLIDE[1]); set_text(ssh, SLIDE[2]); set_paras(bsh, SLIDE[4])
s.part.related_part(psh._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = (SH / SLIDE[3]).read_bytes()
prs.save(DECK); print(DECK, "слайдов:", len(prs.slides))
