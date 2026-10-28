# -*- coding: utf-8 -*-
"""v1.48.0 deliverables: board tools (ergonomics, live data, templates).
Bumps version strings, adds board-tools text to the user guide / descriptions / deck, refreshes board screenshots.
Idempotent: text is inserted only once; version replaced only if old token present."""
import glob, os, copy, pathlib, docx
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
os.chdir(pathlib.Path(__file__).resolve().parents[2])
OLDV, NEWV = "1.47.0", "1.48.0"
SHOT = pathlib.Path("build/shots/board.jpg").read_bytes()
MARK = "v1.48"
NOTE = ("Новое в v1.48: шаблоны ритуалов одной кнопкой (ретро Start/Stop/Continue, 4L, Mad/Sad/Glad, «Кораблик», "
        "Lean Coffee, SWOT, «Влияние/Усилия», Premortem, Customer Journey Map, Lean Canvas, канбан-колонки); "
        "живые карточки задач из бэклога и виджеты-метрики; превращение стикеров в задачи, риски, импедименты, "
        "решения (ADR), action items и уроки; отмена/повтор, копирование, группировка, закрепление, выравнивание, "
        "привязка к сетке, мини-карта и поиск по доске.")
UG_58 = [
 "Шаблоны (📋): ретро Start/Stop/Continue, 4L, Mad/Sad/Glad, «Кораблик», Lean Coffee, SWOT, матрица «Влияние/Усилия», Premortem, Customer Journey Map, Lean Canvas и канбан-колонки. Шаблон вставляется в центр экрана; клик по рамке выделяет его целиком, и он двигается вместе со стикерами.",
 "Данные (📊): живые карточки задач из бэклога (статус, исполнитель, SP; клик по статусу меняет его, двойной клик открывает задачу в бэклоге), кнопка «Спринт на доску» раскладывает задачи текущего спринта по колонкам статусов, виджеты-метрики (velocity, выполнение и цель спринта, health-index, риск срыва, открытые баги) пересчитываются автоматически.",
 "Стикер → сущность: выделите стикеры и нажмите «⇢ В сущность» — они станут задачами (стикер превращается в живую карточку), рисками, импедиментами, решениями (ADR), action items ретроспективы или уроками; на стикере появится метка-ссылка на созданную запись.",
 "Работа с выделенным: панель действий внизу — дублировать, сгруппировать/разгруппировать, закрепить (закреплённое нельзя сдвинуть и удалить), выровнять по краям и центру, распределить равномерно. Фрейм переносится вместе с содержимым.",
 "Отмена и повтор (Ctrl+Z / Ctrl+Shift+Z или Ctrl+Y, кнопки ↶ ↷), копировать/вставить/дублировать (Ctrl+C/V/D), группа (Ctrl+G), закрепить (Ctrl+L), поиск по тексту на доске (Ctrl+F, Enter — следующее совпадение), привязка к сетке (⌗) и мини-карта (🗺): клик по ней переносит к нужной области.",
]
UG_1715 = [
 "«↶ / ↷» — отмена и повтор действий на доске. Как работает: история хранит последние 50 шагов; при изменениях соавторов в реальном времени локальная история начинается заново, чтобы отмена не откатила чужую работу.",
 "«🔍 Поиск» — поиск по тексту стикеров, фигур, ссылок и карточек задач. Как работает: совпадения выделяются; Enter переводит камеру к следующему.",
 "«📋 Шаблоны» и «📊 Данные» — меню шаблонов ритуалов и живых данных приложения (карточки задач, спринт на доске, виджеты-метрики).",
 "«⌗» — привязка к сетке; «🗺» — показать/скрыть мини-карту.",
 "Панель действий (появляется при выделении) — ⧉ дублировать, ▣/⊟ группа, 🔒 закрепить, ⇤ ⇔ ⇥ ⤒ ⇕ ⤓ выравнивание, ⇹ ⇳ распределение, «⇢ В сущность», 🗑 удалить.",
]

def fixruns(p):
    for r in p.runs:
        if OLDV in r.text: r.text = r.text.replace(OLDV, NEWV)

def all_pars(d):
    for p in d.paragraphs: yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: yield p
    for s in d.sections:
        for hf in (s.header, s.footer):
            for p in hf.paragraphs: yield p

def next_heading_idx(ps, i):
    for j in range(i + 1, len(ps)):
        if ps[j].style.name.startswith("Heading"): return j
    return None

def replace_board_images(d):
    cur, n = "", 0
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"): cur = p.text; continue
        if "Доска" not in cur: continue
        for bl in p._p.findall('.//' + qn('a:blip')):
            d.part.related_parts[bl.get(qn('r:embed'))]._blob = SHOT; n += 1
    return n

for fn in sorted(glob.glob("*.docx")) + sorted(glob.glob("Тестовая документация/*.docx")):
    if "~$" in fn: continue
    d = docx.Document(fn)
    for p in all_pars(d): fixruns(p)
    ps = d.paragraphs; added = imgs = 0
    already = any(MARK in p.text and "шаблоны ритуалов" in p.text for p in ps) or any(p.text.startswith("Шаблоны (📋)") for p in ps)
    if fn.startswith("Руководство_пользователя") and not already:
        for title, lines in (("5.8 Доска", UG_58), ("17.15 Доска", UG_1715)):
            i = next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and p.text.startswith(title)), None)
            if i is None: continue
            j = next_heading_idx(ps, i)
            for ln in lines:
                if j is not None: ps[j].insert_paragraph_before(ln, style="List Bullet")
                else: d.add_paragraph(ln, style="List Bullet")
                added += 1
            ps = d.paragraphs
    if (fn.startswith("Иллюстрированное_описание") or fn.startswith("Описание_")) and not already:
        i = next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and p.text.startswith("Доска")), None)
        if i is not None:
            j = next_heading_idx(ps, i)
            ref = next((p for p in ps[i + 1:j] if p.text.startswith("Новое в v1.41")), ps[i + 1])
            if j is not None: np_ = ps[j].insert_paragraph_before(NOTE)
            else: np_ = d.add_paragraph(NOTE)
            np_.style = ref.style; added += 1
    if fn.startswith("Руководство_пользователя") or fn.startswith("Иллюстрированное_описание"):
        imgs = replace_board_images(d)
    d.save(fn); print("docx %-60s +%d text, %d images" % (fn[:60], added, imgs))

prs = Presentation("Презентация_приложения.pptx")
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs: fixruns(p)
    texts = [sh for sh in s.shapes if sh.has_text_frame]
    if not any("Доска — интерактивный бесконечный холст" in sh.text_frame.text for sh in texts): continue
    body = next(sh for sh in texts if sh.text_frame.text.strip().startswith("•"))
    if MARK not in body.text_frame.text:
        last = body.text_frame.paragraphs[-1]._p
        for line in ("•  Шаблоны (v1.48): ретро, SWOT, «Влияние/Усилия», Premortem, CJM, Lean Canvas — одной кнопкой.",
                     "•  Живые карточки задач и виджеты-метрики; стикер → задача/риск/решение; отмена, группы, мини-карта."):
            np_ = copy.deepcopy(last); last.addnext(np_); last = np_
            runs = np_.findall(qn('a:r'))
            for r in runs[1:]: np_.remove(r)
            runs[0].find(qn('a:t')).text = line
    for sh in s.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            bl = sh._element.find('.//' + qn('a:blip')); s.part.related_part(bl.get(qn('r:embed')))._blob = SHOT
    print("deck: board slide updated")
prs.save("Презентация_приложения.pptx"); print("done")
