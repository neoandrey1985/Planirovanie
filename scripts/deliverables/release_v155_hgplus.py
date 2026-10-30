# -*- coding: utf-8 -*-
"""v1.55.0: «Гант по иерархии» — масштаб, вехи и релизы, фильтры, колонки, перетаскивание, базовый план, прогноз,
загрузка, экспорт, критический путь — в документах и презентации; миграция V15 (tasks.baseline) в руководстве
администратора и инструкции; версия во всех .docx. Снимки — capture_special_shots.py (s_hgweek, s_hgplan, s_hgfilter).
Идемпотентно. Запуск: python scripts/deliverables/release_v155_hgplus.py"""
import copy, glob, os, pathlib, docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.54.1", "1.55.0"
MIG, PREV = "V15", "V14"
ADM_LINE = "v1.55 (миграция V15): tasks.baseline — базовый план «Ганта по иерархии» (JSON: a — начало, b — конец, at — дата фиксации). Входит в снимок состояния /api/state."
DEP = [("миграции V1–V14)", "миграции V1–V15)"),
       ("V14 (tasks.stages и params.stage_roles — метки этапов).", "V14 (tasks.stages и params.stage_roles — метки этапов), V15 (tasks.baseline — базовый план)."),
       ("При старте Flyway применит миграции V1…V14.", "При старте Flyway применит миграции V1…V15."),
       ("Не редактируйте уже применённые Flyway-миграции (V1…V14) — добавляйте только новые (V15, …).", "Не редактируйте уже применённые Flyway-миграции (V1…V15) — добавляйте только новые (V16, …).")]
DEP_LINE = "v1.55: при старте применяются миграции Flyway V1…V15 (V15 добавляет tasks.baseline — базовый план «Ганта по иерархии»)."
EXTRA_UG = []          # дополнительные разделы руководства (в «Управлении стримами и командами» — «Гант стрима»)
EXTRA_DESC = []
EXTRA_SLIDES = []
UG = ("6. «Гант по иерархии»: масштаб, фильтры, планирование и экспорт", "Над диаграммой три строки инструментов: уровень, масштаб, колонки и экспорт; фильтры и поиск; что показывать и базовый план.",
      ["Масштаб: «Вписать» (вся шкала на экране), «Недели», «Спринты», «Месяцы», «Кварталы». В масштабах, кроме «Вписать», шкала прокручивается по горизонтали, левая колонка закреплена; «К сегодня» прокручивает к текущей дате. Выбор запоминается.",
       "Вехи ⚑ и релизы 🚀 — подписями во второй строке шапки и линиями по всей диаграмме (конец спринта вехи или последнего спринта релиза).",
       "Фильтры и поиск: по ID или названию, исполнителю («Только мои» — по «Моему профилю»), роли, статусу, релизу, спринту; «просроченные», «конфликты», «только критический путь». Найденные элементы показываются вместе с предками, «Сбросить» возвращает всё.",
       "Колонки слева (меню «Колонки»): исполнитель, статус, спринт, плановые даты, остаток SP, резерв — сколько спринтов задачу можно сдвинуть, не задев преемника (отрицательный — конфликт).",
       "Перетаскивание: редактор тянет плановую полосу задачи мышью в другой спринт; подсказка показывает новый спринт и конфликты (раньше предшественника, позже преемника, за спринтом родителя); после переноса — «Отменить».",
       "Базовый план: «Зафиксировать план» сохраняет текущие сроки всех элементов; затем исходный план виден пунктирной рамкой, сдвиг — «+1 спр.» у задачи и «сдвиг …» у эпиков и фич, KPI «Сдвиг сроков» считает сдвинутые элементы. «Сбросить базовый план» удаляет снимок.",
       "Прогноз: для эпиков и фич с начатой работой — дата завершения по фактическому темпу (готово SP за дни с начала); если позже плана — красный пунктир после полосы.",
       "Загрузка: строка под диаграммой — план против ёмкости по каждому спринту (красный — перегруз), в подсказке разбивка плана по ролям.",
       "Экспорт (меню «Экспорт»): PNG и PDF — снимок диаграммы со стрелками и метками (PDF требует доступа к CDN), Excel — лист «Гант по иерархии»: иерархия, даты плана и факта, объём и прогресс, базовый план и сдвиг, прогноз и все метки этапов; этот лист есть и в общей выгрузке в Excel."],
      [("s_hgweek.jpg", "Масштаб «Недели», прокрутка и колонки"), ("s_hgplan.jpg", "Базовый план, сдвиг, прогноз, вехи и загрузка"), ("s_hgfilter.jpg", "Фильтр по роли с показом предков")])
DESC = ("«Гант по иерархии»: планирование", "Масштаб от недель до кварталов с прокруткой и кнопкой «К сегодня», вехи и релизы на шкале, фильтры и поиск, колонки "
        "(исполнитель, статус, спринт, даты, остаток, резерв), перетаскивание задач между спринтами, базовый план и сдвиг сроков, прогноз завершения "
        "по фактическому темпу, загрузка спринтов и экспорт в PNG, PDF и Excel.")
SLIDE = ("🧭", "Гант: масштаб и планирование", "Неделя, спринт, месяц, квартал.", "s_hgplan.jpg", [
    "•  Вехи и релизы на шкале, фильтры, колонки.",
    "•  Перетаскивание задач между спринтами.",
    "•  Базовый план, сдвиг сроков, прогноз.",
    "•  Загрузка спринтов, экспорт PNG / PDF / Excel."])
ANCHOR_SLIDE = "Зависимости на Ганте"


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def set_par(p, text):
    for r in p.runs[1:]: r.text = ""
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)


def replace_in(d, a, b):
    n = 0
    for p in d.paragraphs:
        if a in p.text and b not in p.text: set_par(p, p.text.replace(a, b)); n += 1
    return n


def after(p, text):
    np_ = copy.deepcopy(p._p); p._p.addnext(np_); q = Paragraph(np_, p._parent); set_par(q, text); return q


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def next_heading(d, i):
    return next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))


def heading(d, text, lvl="Heading 2"):
    return next(k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.strip() == text)


def add_section(d, after_title, sec, width):
    title, text, bullets, pics = sec
    if any(p.text == title for p in d.paragraphs): return False
    anchor = next_heading(d, heading(d, after_title))
    put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
    for b in bullets: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
    for img, cap in pics: add_pic(d, anchor, img, cap, width)
    return True


fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
st = add_section(d, "6. Зависимости на «Ганте по иерархии»", UG, 6.3)
for after_title, sec in EXTRA_UG: st = add_section(d, after_title, sec, 6.3) or st
d.save(fn); print(fn[:58], "разделы:", "добавлены" if st else "уже есть")
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn); n = 0
    for after_title, (title, text, img) in [("Зависимости на Ганте", (DESC[0], DESC[1], "s_hgplan.jpg"))] + EXTRA_DESC:
        if any(p.style.name.startswith("Heading") and p.text.strip() == title for p in d.paragraphs): continue
        anchor = next_heading(d, heading(d, after_title))
        put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
        if pics: add_pic(d, anchor, img, title, 6.61)
        n += 1
    d.save(fn); print(fn[:58], "разделов:", n)
fn = f"Руководство_администратора_{P}.docx"; d = docx.Document(fn); n = 0
n += replace_in(d, f"(миграции V1…{PREV})", f"(миграции V1…{MIG})")
for p in list(d.paragraphs):
    if p.text.startswith("Основная БД") and f"{MIG} — базовый план" not in p.text:
        set_par(p, p.text.replace(f"Flyway V1…{PREV}:", f"Flyway V1…{MIG}:").rstrip(".") + f", {MIG} — базовый план «Ганта по иерархии» (tasks.baseline)."); n += 1
if not any(p.text.startswith(ADM_LINE[:20]) for p in d.paragraphs):
    last = [p for p in d.paragraphs if p.text.startswith("v") and "(миграция V" in p.text][-1]; after(last, ADM_LINE); n += 1
d.save(fn); print(fn[:58], "правок:", n)
fn = f"Инструкция_по_развёртыванию_{P}.docx"; d = docx.Document(fn); n = 0
for a, b in DEP: n += replace_in(d, a, b)
if not any(p.text == DEP_LINE for p in d.paragraphs):
    last = [p for p in d.paragraphs if p.text.startswith("v") and "миграции Flyway" in p.text][-1]; after(last, DEP_LINE); n += 1
d.save(fn); print(fn[:58], "правок:", n)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


DECK = "Презентация_приложения.pptx"; prs = Presentation(DECK)
for anchor_title, sl in [(ANCHOR_SLIDE, SLIDE)] + EXTRA_SLIDES:
    s = next((x for x in prs.slides if title_of(x) == sl[1]), None)
    if s is None:
        i = next(k for k, x in enumerate(prs.slides) if title_of(x) == anchor_title)
        s = clone_after(prs, list(prs.slides)[i], i + 1, sl[3])
    icon, tsh, ssh, bsh, psh = parts(s)
    set_text(icon, sl[0]); set_text(tsh, sl[1]); set_text(ssh, sl[2]); set_paras(bsh, sl[4])
    s.part.related_part(psh._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = (SH / sl[3]).read_bytes()
prs.save(DECK); print(DECK, "слайдов:", len(prs.slides))
