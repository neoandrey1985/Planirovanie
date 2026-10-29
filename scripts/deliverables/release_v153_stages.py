# -*- coding: utf-8 -*-
"""v1.53.0: метки этапов на «Ганте по иерархии» в документах и презентации. Руководство пользователя — раздел
«6. Метки этапов» после «6. Гант по иерархии»; описания — раздел «Метки этапов»; руководство администратора и
инструкция по развёртыванию — миграция V14 (tasks.stages, params.stage_roles); слайд «Метки этапов» после слайда
«Гант по иерархии»; версия во всех .docx. Снимки — capture_special_shots.py (s_stages, s_stlegend, s_stpanel).
Идемпотентно. Запуск: python scripts/deliverables/release_v153_stages.py"""
import copy, glob, os, pathlib, docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.52.0", "1.53.0"
STREAMS = False
MIG, PREV = "V14", "V13"
H = "Метки этапов"
UG = ("6. Метки этапов на «Ганте по иерархии»", "На строках «Ганта по иерархии» — метки этапов от бизнес-анализа до выхода в Прод. Контур фигуры — начало этапа, заливка — завершение; точка в углу — метка поставлена вручную.",
      ["Бизнес-анализ — круг, системный анализ — треугольник, разработка — квадрат, ИФТ (интеграционно-функциональное тестирование) — звезда, ПСИ (приёмо-сдаточные испытания) — ромб, выход в Прод — ракета. Все 12 обозначений с описанием — в легенде под диаграммой.",
       "Автоматически: бизнес-анализ, системный анализ, разработка и ИФТ — по ролям задач (начало — дата «начато», завершение — «готово»). Соответствие ролей этапам задаётся в «Параметрах» → «Метки этапов» (по умолчанию: БА — Product Owner, СА — Analyst, разработка — Backend, Frontend, Designer, DevOps, ИФТ — QA).",
       "ПСИ — по демо спринта готовой задачи (итог «Принято» — завершение). Прод — по релизу задачи: начало выхода в Прод — начало последнего спринта релиза, выход в Прод — его окончание, когда релиз «Готово».",
       "У родителя (эпик, фича, история с подзадачами) метки собираются из его собственных дат и потомков: начало — самое раннее, завершение — когда этап завершён везде, где он начат.",
       "Вручную: клик по названию элемента или по метке — панель «Этапы» с 12 датами (начало и завершение каждого этапа), «сегодня» и «Сбросить всё на авто». Заполненная дата важнее автоматической; пустая — снова автоматическая. Править может редактор, просмотр — только чтение.",
       "Флажок «метки этапов» над диаграммой скрывает и показывает метки. С сервером ручные даты хранятся у задачи."],
      [("s_stages.jpg", "Метки этапов на «Ганте по иерархии»"), ("s_stlegend.jpg", "Легенда меток этапов"), ("s_stpanel.jpg", "Панель «Этапы»: ручные и автоматические даты")])
DESC = (H, "Метки этапов на «Ганте по иерархии»: бизнес-анализ (круг), системный анализ (треугольник), разработка (квадрат), ИФТ (звезда), ПСИ (ромб) "
           "и выход в Прод (ракета); контур — начало, заливка — завершение. Ставятся автоматически — по ролям задач, демо спринта и релизам "
           "(соответствие ролей — в «Параметрах») — или вручную в панели «Этапы»; у родителя собираются из потомков. Все обозначения — в легенде.")
SLIDE = ("🏷", H, "От бизнес-анализа до выхода в Прод.", "s_stages.jpg", [
    "•  Круг, треугольник, квадрат, звезда, ромб, ракета.",
    "•  Контур — начало этапа, заливка — завершение.",
    "•  Авто — по ролям задач, демо и релизам; вручную — панель «Этапы».",
    "•  У эпиков и фич — сводно по потомкам; легенда со всеми метками."])
ADM_LINE = f"v1.53 (миграция {MIG}): tasks.stages — ручные даты меток этапов задачи (JSON: ba_s, ba_e, … prod_e) и params.stage_roles — роли этапов для автоматических меток. Входят в снимок состояния /api/state."
ADM_TAIL = f", {MIG} — метки этапов «Ганта по иерархии» (tasks.stages, params.stage_roles)."
DEP = [("миграции V1–V13)", "миграции V1–V14)"),
       ("V13 (app_users.member_name — «Мой профиль»).", "V13 (app_users.member_name — «Мой профиль»), V14 (tasks.stages и params.stage_roles — метки этапов)."),
       ("При старте Flyway применит миграции V1…V13.", "При старте Flyway применит миграции V1…V14."),
       ("Не редактируйте уже применённые Flyway-миграции (V1…V13) — добавляйте только новые (V14, …).", "Не редактируйте уже применённые Flyway-миграции (V1…V14) — добавляйте только новые (V15, …).")]
DEP_LINE = "v1.53: при старте применяются миграции Flyway V1…V14 (V14 добавляет tasks.stages и params.stage_roles для меток этапов)."
STREAM_NOTE = None


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def set_par(p, text):
    for r in p.runs[1:]: r.text = ""
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)


def replace_in(d, a, b):
    n = 0
    for p in d.paragraphs:
        if a in p.text and b not in p.text: set_par(p, p.text.replace(a, b)); n += 1
    return n


def after(p, text):
    np_ = copy.deepcopy(p._p); p._p.addnext(np_); q = Paragraph(np_, p._parent); set_par(q, text); return q


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def next_heading(d, i):
    return next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))


def heading(d, text, lvl="Heading 2"):
    return next(k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.strip() == text)


# ---------- руководство пользователя ----------
fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
if not any(p.text == UG[0] for p in d.paragraphs):
    if STREAM_NOTE:
        i = heading(d, "6. Гант по иерархии"); k = next(x for x in range(i + 1, len(d.paragraphs)) if d.paragraphs[x].style.name == "List Bullet")
        d.paragraphs[k].insert_paragraph_before(STREAM_NOTE, style="List Bullet")
    anchor = next_heading(d, heading(d, "6. Гант по иерархии"))
    put_before(anchor, d.add_paragraph(UG[0], style="Heading 2")); put_before(anchor, d.add_paragraph(UG[1]))
    for b in UG[2]: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
    for img, cap in UG[3]: add_pic(d, anchor, img, cap, 6.3)
    st = "добавлен"
else: st = "уже есть"
d.save(fn); print(fn[:58], "раздел:", st)
# ---------- описания ----------
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not any(p.style.name.startswith("Heading") and p.text.strip() == H for p in d.paragraphs):
        anchor = next_heading(d, heading(d, "Гант по иерархии"))
        put_before(anchor, d.add_paragraph(H, style="Heading 2")); put_before(anchor, d.add_paragraph(DESC[1]))
        if pics: add_pic(d, anchor, "s_stages.jpg", H, 6.61)
        st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "раздел:", st)
# ---------- руководство администратора ----------
fn = f"Руководство_администратора_{P}.docx"; d = docx.Document(fn); n = 0
n += replace_in(d, f"(миграции V1…{PREV})", f"(миграции V1…{MIG})")
for p in list(d.paragraphs):
    if p.text.startswith("Основная БД") and f"{MIG} — метки этапов" not in p.text:
        set_par(p, p.text.replace(f"Flyway V1…{PREV}:", f"Flyway V1…{MIG}:").rstrip(".") + ADM_TAIL); n += 1
if not any(p.text.startswith(ADM_LINE[:20]) for p in d.paragraphs):
    last = [p for p in d.paragraphs if p.text.startswith("v1.") and "(миграция V" in p.text][-1]; after(last, ADM_LINE); n += 1
d.save(fn); print(fn[:58], "правок:", n)
# ---------- инструкция по развёртыванию ----------
fn = f"Инструкция_по_развёртыванию_{P}.docx"; d = docx.Document(fn); n = 0
for a, b in DEP: n += replace_in(d, a, b)
if DEP_LINE and not any(p.text == DEP_LINE for p in d.paragraphs):
    last = [p for p in d.paragraphs if p.text.startswith("v1.") and "миграции Flyway" in p.text][-1]; after(last, DEP_LINE); n += 1
d.save(fn); print(fn[:58], "правок:", n)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")


# ---------- слайд ----------
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


DECK = "Презентация_приложения.pptx"; prs = Presentation(DECK)
s = next((x for x in prs.slides if title_of(x) == SLIDE[1]), None)
if s is None:
    i = next(k for k, x in enumerate(prs.slides) if title_of(x) == "Гант по иерархии")
    s = clone_after(prs, list(prs.slides)[i], i + 1, SLIDE[3])
icon, tsh, ssh, bsh, psh = parts(s)
set_text(icon, SLIDE[0]); set_text(tsh, SLIDE[1]); set_text(ssh, SLIDE[2]); set_paras(bsh, SLIDE[4])
s.part.related_part(psh._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = (SH / SLIDE[3]).read_bytes()
prs.save(DECK); print(DECK, "слайдов:", len(prs.slides))
