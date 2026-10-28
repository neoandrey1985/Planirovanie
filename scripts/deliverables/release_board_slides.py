# -*- coding: utf-8 -*-
"""Отдельные слайды про «Доску» в Презентация_приложения.pptx.
Клонирует оформление обзорного слайда «Доска — интерактивный бесконечный холст» (иконка, заголовок,
подзаголовок, пункты, снимок) и вставляет сразу после него пять слайдов. Снимки — из
capture_board_shots.py (build/shots/board*.jpg). Идемпотентно: существующие слайды обновляются, не дублируются.
Запуск: python scripts/deliverables/release_board_slides.py"""
import copy, pathlib, os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
os.chdir(pathlib.Path(__file__).resolve().parents[2])
DECK = "Презентация_приложения.pptx"
SHOTS = pathlib.Path("build/shots")
OVERVIEW = "Доска — интерактивный бесконечный холст"
SLIDES = [
 ("🧰", "Доска — эргономика холста", "Работа с холстом как в профессиональных редакторах.", "board_ergo.jpg", [
   "•  Отмена и повтор (Ctrl+Z / Ctrl+Y), копировать, вставить и дублировать (Ctrl+C/V/D) — связи между копиями сохраняются.",
   "•  Группы, закрепление 🔒, выравнивание и равномерное распределение, привязка к сетке.",
   "•  Мини-карта и поиск по тексту (Ctrl+F); фрейм переносится вместе с содержимым.",
   "•  Панель действий появляется при выделении — всё нужное под рукой."]),
 ("📋", "Доска — шаблоны ритуалов", "Готовая структура встречи одной кнопкой.", "board_tpl.jpg", [
   "•  Ретро: Start/Stop/Continue, 4L, Mad/Sad/Glad, «Кораблик».",
   "•  Анализ: Lean Coffee, SWOT, матрица «Влияние/Усилия», Premortem.",
   "•  Продукт и процесс: Customer Journey Map, Lean Canvas, канбан-колонки, дорожки ответственности.",
   "•  Шаблон двигается целиком вместе со стикерами; цвет стикера подсказывает колонку."]),
 ("📊", "Доска — живые данные приложения", "Холст связан с бэклогом и метриками.", "board.jpg", [
   "•  Карточки задач из бэклога: статус меняется кликом, двойной клик открывает задачу.",
   "•  «Спринт на доску»: задачи текущего спринта по колонкам статусов.",
   "•  Виджеты-метрики: velocity, выполнение и цель спринта, health-index, риск срыва, баги.",
   "•  Стикер превращается в задачу, риск, импедимент, решение (ADR), action item или урок."]),
 ("🗳️", "Доска — проведение сессий", "Фасилитация ретро и воркшопов.", "board_sess.jpg", [
   "•  Голосование точками: лимит на участника, скрытый режим, раскрытие итогов и медали лидеров.",
   "•  Анонимный ввод: авторы стикеров скрыты до раскрытия.",
   "•  Показ по фреймам как по слайдам и режим «следовать за ведущим».",
   "•  Комментарии к элементам с @упоминаниями и сводка обсуждений."]),
 ("🔷", "Доска — схемы и диаграммы", "Процессы, идеи и данные на одном холсте.", "board2.jpg", [
   "•  Блок-схемы: начало/конец, процесс, решение, ввод/вывод, документ, база данных; подписи стрелок.",
   "•  Майнд-карта: Tab — ветка, Enter — соседняя, авто-раскладка в обе стороны.",
   "•  Таблица на холсте и дорожки ответственности по ролям команды.",
   "•  Авто-раскладка выделенного сеткой или деревом по связям."]),
]
OVERVIEW_TAIL = "•  Подробнее — на следующих пяти слайдах: эргономика, шаблоны, живые данные, сессии и схемы."

def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else ""
def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]  # обёртки python-pptx каждый раз новые — сравниваем XML-элементы
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic           # иконка, заголовок, подзаголовок, пункты, снимок

def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text

def set_bullets(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)

def set_picture(slide, pic, path):
    slide.part.related_part(pic._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = pathlib.Path(path).read_bytes()

def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    new.shapes.add_picture(str(SHOTS / shot), pic.left, pic.top, pic.width, pic.height)  # свой снимок: одинаковые картинки в pptx общие
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new

prs = Presentation(DECK)
slides = list(prs.slides)
src_i = next(i for i, s in enumerate(slides) if title_of(s) == OVERVIEW); src = slides[src_i]
_, _, _, body, pic = parts(src)
set_picture(src, pic, SHOTS / "board.jpg")
keep = [p.text for p in body.text_frame.paragraphs if not any(k in p.text for k in ("v1.48", "v1.49", "Живые карточки задач и виджеты", "Подробнее — на следующих"))]
set_bullets(body, keep + [OVERVIEW_TAIL])
for k, (icon_t, title, sub, shot, bullets) in enumerate(SLIDES):
    cur = list(prs.slides); s = next((x for x in cur if title_of(x) == title), None)
    if s is None: s = clone_after(prs, src, src_i + 1 + k, shot)
    icon, tsh, ssh, bsh, psh = parts(s)
    set_text(icon, icon_t); set_text(tsh, title); set_text(ssh, sub); set_bullets(bsh, bullets); set_picture(s, psh, SHOTS / shot)
prs.save(DECK)
for i, s in enumerate(prs.slides):
    if "Доска" in title_of(s) or i == len(prs.slides) - 1: print(i, title_of(s) or texts(s)[0].text_frame.text[:30])
