# -*- coding: utf-8 -*-
"""v1.56.0: «Гант по иерархии» стал легче — режимы вида, одна строка инструментов, показатели строкой, легенда по кнопке,
спокойная графика, фокус на элементе, плотность строк, выгрузка в PPTX по эпикам — в документах и презентации;
версия во всех .docx. Снимки — capture_special_shots.py (s_hglight, s_hgfocus, s_hgmenu). Миграций нет.
Идемпотентно. Запуск: python scripts/deliverables/release_v156_hglight.py"""
import copy, glob, os, pathlib, docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.55.0", "1.56.0"
# устаревшие формулировки раздела «6. «Гант по иерархии»: масштаб, фильтры…» после облегчения интерфейса
REPL = [("Над диаграммой три строки инструментов: уровень, масштаб, колонки и экспорт; фильтры и поиск; что показывать и базовый план.",
         "Над диаграммой одна строка инструментов: уровень, режим вида, «Настроить», масштаб, поиск, «Фильтры», меню «⋯» (колонки, плотность, базовый план, экспорт) и «Легенда»."),
        ("Колонки слева (меню «Колонки»)", "Колонки слева (меню «⋯» → «Колонки»)"),
        ("Экспорт (меню «Экспорт»): PNG и PDF", "Экспорт (меню «⋯» → «Экспорт»): PPTX — презентация по эпикам, PNG и PDF")]
EXTRA_UG = []
EXTRA_DESC = []
EXTRA_SLIDES = []
UG = ("6. «Гант по иерархии»: лёгкий вид, фокус и презентация", "Страница по умолчанию открывается в спокойном виде: только сводные полосы, прогресс и отклонения; подробности включаются по необходимости.",
      ["Режимы вида: «Обзор» (по умолчанию — полосы, прогресс, вехи), «План» (+ зависимости и загрузка спринтов), «Контроль» (+ факт, базовый план и прогноз), «Этапы» (+ метки этапов). "
       "Отдельные слои включаются в «Настроить» — вид становится «Свой вид». Выбор запоминается.",
       "Одна строка инструментов: уровень (Эпики, До фич, До историй, Всё), режим, «Настроить», масштаб, поиск, «Фильтры» (с числом активных), меню «⋯» — колонки, плотность строк, базовый план и экспорт — и «Легенда».",
       "Показатели одной строкой: состав и объём, готовность, «просрочено», «без родителя», «сдвиг сроков». Клик по жёлтой или красной плашке включает соответствующий фильтр, повторный — снимает.",
       "Легенда по кнопке «Легенда» и только для включённых слоёв: метки этапов — в режиме «Этапы», стрелки — когда включены зависимости.",
       "Спокойная графика: тонкие полосы, светлая сетка, готовые задачи нейтрального цвета — цветом выделяются только просрочка, конфликты, прогноз и сдвиг; процент готовности — мини-полоской рядом с объёмом.",
       "Фокус на элементе: клик по строке или полосе оставляет яркими её предков, потомков и связанные зависимостями элементы, стрелки — только у этой цепочки; плашка «фокус: …» и Esc снимают фокус.",
       "Плотность строк: «⋯» → «компактно» — больше строк на экране.",
       "Презентация PPTX по эпикам: «⋯» → «PPTX — презентация по эпикам» — титул, обзор эпиков и фич, слайд на каждый эпик (до задач, со стрелками и вехами) и сводка «Что требует внимания» — просрочено, сдвиг, прогноз позже плана, конфликты зависимостей."],
      [("s_hglight.jpg", "Режим «Обзор»: показатели строкой, одна строка инструментов, спокойные полосы"), ("s_hgfocus.jpg", "Фокус на T-14: цепочка предков, потомков и зависимостей"), ("s_hgmenu.jpg", "Режим «Контроль» и меню «⋯»: колонки, плотность, базовый план, экспорт в PPTX")])
DESC = ("«Гант по иерархии»: лёгкий вид и презентация", "Режимы вида «Обзор / План / Контроль / Этапы», одна строка инструментов, показатели-фильтры строкой, легенда по кнопке, "
        "спокойная графика с акцентом на отклонениях, фокус на цепочке элемента, компактные строки и выгрузка в PowerPoint: обзор, слайд на каждый эпик и сводка «что требует внимания».")
SLIDE = ("🪶", "Гант: лёгкий вид и PPTX", "Режимы, фокус, презентация.", "s_hgfocus.jpg", [
    "•  Обзор / План / Контроль / Этапы.",
    "•  Одна строка инструментов, легенда по кнопке.",
    "•  Фокус на цепочке элемента, компактные строки.",
    "•  PPTX: обзор, слайд на эпик, сводка."])
ANCHOR_SLIDE = "Гант: масштаб и планирование"


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def set_par(p, text):
    for r in p.runs[1:]: r.text = ""
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)


def replace_in(d, a, b):
    n = 0
    for p in d.paragraphs:
        if a in p.text and b not in p.text: set_par(p, p.text.replace(a, b)); n += 1
    return n


def after(p, text):
    np_ = copy.deepcopy(p._p); p._p.addnext(np_); q = Paragraph(np_, p._parent); set_par(q, text); return q


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def next_heading(d, i):
    return next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))


def heading(d, text, lvl="Heading 2"):
    return next(k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.strip() == text)


def add_section(d, after_title, sec, width):
    title, text, bullets, pics = sec
    if any(p.text == title for p in d.paragraphs): return False
    anchor = next_heading(d, heading(d, after_title))
    put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
    for b in bullets: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
    for img, cap in pics: add_pic(d, anchor, img, cap, width)
    return True


fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
for a, b in REPL: replace_in(d, a, b)
st = add_section(d, "6. «Гант по иерархии»: масштаб, фильтры, планирование и экспорт", UG, 6.3)
for after_title, sec in EXTRA_UG: st = add_section(d, after_title, sec, 6.3) or st
d.save(fn); print(fn[:58], "разделы:", "добавлены" if st else "уже есть")
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn); n = 0
    for after_title, (title, text, img) in [("«Гант по иерархии»: планирование", (DESC[0], DESC[1], "s_hglight.jpg"))] + EXTRA_DESC:
        if any(p.style.name.startswith("Heading") and p.text.strip() == title for p in d.paragraphs): continue
        anchor = next_heading(d, heading(d, after_title))
        put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
        if pics: add_pic(d, anchor, img, title, 6.61)
        n += 1
    d.save(fn); print(fn[:58], "разделов:", n)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


DECK = "Презентация_приложения.pptx"; prs = Presentation(DECK)
for anchor_title, sl in [(ANCHOR_SLIDE, SLIDE)] + EXTRA_SLIDES:
    s = next((x for x in prs.slides if title_of(x) == sl[1]), None)
    if s is None:
        i = next(k for k, x in enumerate(prs.slides) if title_of(x) == anchor_title)
        s = clone_after(prs, list(prs.slides)[i], i + 1, sl[3])
    icon, tsh, ssh, bsh, psh = parts(s)
    set_text(icon, sl[0]); set_text(tsh, sl[1]); set_text(ssh, sl[2]); set_paras(bsh, sl[4])
    s.part.related_part(psh._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = (SH / sl[3]).read_bytes()
prs.save(DECK); print(DECK, "слайдов:", len(prs.slides))
