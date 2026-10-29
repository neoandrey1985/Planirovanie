# -*- coding: utf-8 -*-
"""v1.52.0: «Гант по иерархии» в документах и презентации. Руководство пользователя — раздел «6. Гант по иерархии»
после «6. Гант» и ссылка из «18. Иерархия работ (WBS)»; описания — раздел после «Диаграммы Ганта»; слайд после
слайда «Диаграмма Ганта» (оформление клонируется с него); версия во всех .docx. Снимок — capture_special_shots.py
(s_hgantt). Идемпотентно. Запуск: python scripts/deliverables/release_v152_hgantt.py"""
import copy, glob, os, pathlib, docx
from docx.shared import Inches
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.51.0", "1.52.0"
H = "Гант по иерархии"
UG = ("6. Гант по иерархии", "Вкладка «Расписания» рядом с «Гантом»: Эпик → Фича → История → Задача по полю «Родитель» на шкале спринтов.",
      ["Строка родителя — сводная полоса от самого раннего начала до самого позднего окончания всех потомков; заливка — прогресс по объёму (готовые SP / все SP). Объём и прогресс суммируются снизу вверх, как в «Иерархии работ».",
       "План — по спринтам (полоса), факт — тонкая полоса по датам «начато / готово» (незавершённая — до «сегодня»); красная вертикаль — «сегодня».",
       "Красная рамка — просрочено (спринт прошёл, а элемент не готов) или потомки выходят за спринт родителя; справа от сводной полосы — процент и число просроченных внутри.",
       "Уровни «Эпики / До фич / До историй / Всё» (по умолчанию — до фич), стрелка ▸ раскрывает любой элемент; флажок «факт» скрывает фактические полосы.",
       "Клик по ID открывает элемент в «Бэклоге»; иерархия задаётся полем «Родитель» (Бэклог). Если поле образует цикл, раздел предупредит и покажет остальное.",
       "KPI: эпики и фичи, общий объём и прогресс, просроченные, задачи без родителя."])
WBS_NOTE = "Та же иерархия на шкале времени — вкладка «Гант по иерархии» в «Расписании» (раздел 6)."
DESC = (H, "Эпик → Фича → История → Задача по полю «Родитель» на шкале спринтов: сводные полосы родителей (от раннего начала до позднего "
           "окончания потомков) с заливкой прогресса по объёму, план по спринтам и факт по датам «начато / готово», красная рамка — просрочка; "
           "уровни раскрываются, клик по ID открывает элемент в «Бэклоге».")
SLIDE = ("🪜", H, "Эпик → Фича → История → Задача на шкале.", "s_hgantt.jpg", [
    "•  Сводные полосы: сроки и прогресс по объёму потомков.",
    "•  План по спринтам и факт по датам «начато / готово».",
    "•  Просрочка и выход за спринт родителя — красным.",
    "•  Уровни: эпики, до фич, до историй, всё."])


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def next_heading(d, i):
    return next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))


def heading(d, text, lvl="Heading 2"):
    return next(k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.strip() == text)


# ---------- руководство пользователя ----------
fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
if not any(p.text == UG[0] for p in d.paragraphs):
    anchor = next_heading(d, heading(d, "6. Гант"))
    put_before(anchor, d.add_paragraph(UG[0], style="Heading 2")); put_before(anchor, d.add_paragraph(UG[1]))
    for b in UG[2]: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
    add_pic(d, anchor, "s_hgantt.jpg", H, 6.3)
    i = heading(d, "18. Иерархия работ (WBS)"); next_heading(d, i).insert_paragraph_before(WBS_NOTE)
    st = "добавлен"
else: st = "уже есть"
d.save(fn); print(fn[:58], "раздел:", st)
# ---------- описания ----------
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not any(p.style.name.startswith("Heading") and p.text.strip() == H for p in d.paragraphs):
        anchor = next_heading(d, heading(d, "Диаграмма Ганта"))
        put_before(anchor, d.add_paragraph(H, style="Heading 2")); put_before(anchor, d.add_paragraph(DESC[1]))
        if pics: add_pic(d, anchor, "s_hgantt.jpg", H, 6.61)
        st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "раздел:", st)
for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")


# ---------- слайд ----------
def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


DECK = "Презентация_приложения.pptx"; prs = Presentation(DECK)
s = next((x for x in prs.slides if title_of(x) == SLIDE[1]), None)
if s is None:
    i = next(k for k, x in enumerate(prs.slides) if title_of(x) == "Диаграмма Ганта")
    s = clone_after(prs, list(prs.slides)[i], i + 1, SLIDE[3])
icon, tsh, ssh, bsh, psh = parts(s)
set_text(icon, SLIDE[0]); set_text(tsh, SLIDE[1]); set_text(ssh, SLIDE[2]); set_paras(bsh, SLIDE[4])
s.part.related_part(psh._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = (SH / SLIDE[3]).read_bytes()
prs.save(DECK); print(DECK, "слайдов:", len(prs.slides))
