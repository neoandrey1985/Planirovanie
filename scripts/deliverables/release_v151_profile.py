# -*- coding: utf-8 -*-
"""v1.51.0: «Мой профиль» в документах. Руководство пользователя — глава 21 + правки 20.1 и справочника шапки 17.1;
описания — раздел «Мой профиль»; руководство администратора — эндпоинты /api/auth/me (PUT) и /api/auth/me/password,
миграция V13, события журнала; инструкция по развёртыванию — актуальный список миграций V1…V13; версия во всех .docx.
Слайд — release_v151_slides.py, снимки — capture_special_shots.py (s_profile, s_profile2). Идемпотентно.
Запуск: python scripts/deliverables/release_v151_profile.py"""
import glob, os, pathlib, copy, docx
from docx.shared import Inches

os.chdir(pathlib.Path(__file__).resolve().parents[2])
SH = pathlib.Path("build/shots")
P = "Планирование_спринта_PRO"
OLDV, NEWV = "1.50.0", "1.51.0"
CH = "21. Мой профиль"
SECTIONS = [
 ("21.1 Где открыть", "«👤 Мой профиль» — первый пункт меню пользователя (аватар ▾ в шапке) и вкладка рядом с «Моим днём».", [], []),
 ("21.2 Кто я", "Имя, логин и роль (роль меняет администратор). Поле «Я в команде» связывает вас с участником из раздела «Команда» — по нему строятся «Мой день» и все блоки профиля.",
  ["Без сервера учётной записи нет: привязка и настройки хранятся в этом браузере.",
   "С сервером привязка хранится в учётной записи и работает на любом устройстве; выбор «Я —» на «Моём дне» меняет ту же привязку."],
  [("s_profile.jpg", "Мой профиль")]),
 ("21.3 Моя работа и моё время", "Сводка по вам из данных команды; клик по строке открывает нужный раздел.",
  ["Моя работа: открытые задачи по группам (текущий спринт, следующие, без спринта и прошлые), баги по вашим задачам, ваши импедименты и риски, компетенции из Star Map.",
   "Моё время: загрузка в текущем спринте (план ваших задач к вашей ёмкости с учётом доступности, отсутствий и праздников), отпуска и отсутствия, день рождения."], []),
 ("21.4 Упоминания", "Комментарии на доске, где вас упомянули через @Имя, — последние 10, новые сверху.", [], []),
 ("21.5 Настройки и пароль", "Тема, плотность и стартовая страница переключаются здесь (раньше — в меню пользователя).",
  ["С сервером — смена своего пароля: текущий пароль, новый (не короче 8 символов) и повтор. Доступна любой роли, включая «Просмотр»; после смены другие сеансы этой учётной записи завершаются, текущий остаётся.",
   "Подбор текущего пароля ограничен: после нескольких неверных попыток смена временно блокируется.",
   "«⎋ Выйти» — в профиле и в меню пользователя."],
  [("s_profile2.jpg", "Настройки и смена пароля")]),
]
UG_201_OLD = "Меню пользователя (аватар ▾): роль и вход, данные (Excel, отчёт PPTX, JSON, импорт JSON и CSV/Jira, сброс), тема, плотность, стартовая страница, видеообзор."
UG_201_NEW = "Меню пользователя (аватар ▾): «Мой профиль» (там же тема, плотность, стартовая страница и пароль — глава 21), роль и вход, данные (Excel, отчёт PPTX, JSON, импорт JSON и CSV/Jira, сброс), видеообзор, «Выйти»."
UG_CH2_OLD = "Меню пользователя (аватар ▾): роль и вход, данные (Excel, отчёт PPTX, JSON, импорт, сброс), тема, плотность, стартовая страница, видеообзор."
UG_CH2_NEW = "Меню пользователя (аватар ▾): «Мой профиль» (тема, плотность, стартовая страница, пароль), роль и вход, данные (Excel, отчёт PPTX, JSON, импорт, сброс), видеообзор, «Выйти»."
CTRL = [("«👁 Просмотр / ✎ Редактор / 🛡 Администратор» (первый пункт меню пользователя) — ", "«👁 Просмотр / ✎ Редактор / 🛡 Администратор» (в меню пользователя, раздел «Доступ») — "),
        ("«◐ Светлая / тёмная тема» (в меню пользователя) — ", "«◐ Тема» (в «Моём профиле», блок «Настройки») — "),
        ("«▤ Плотность: комфортно / компактно» (в меню пользователя) — ", "«▤ Плотность: комфортно / компактно» (в «Моём профиле») — "),
        ("«🏠 Стартовая страница» (в меню пользователя) — ", "«🏠 Стартовая страница» (в «Моём профиле») — ")]
CTRL_NEW = [("«👤 Мой профиль» (первый пункт меню пользователя) — ", "открывает профиль: кто я, моя работа и время, упоминания, настройки. ",
             "Как работает: переходит в раздел «Мой профиль» (вкладка рядом с «Моим днём»); подробности — глава 21."),
            ("«⎋ Выйти» (в меню пользователя, при работе с сервером) — ", "завершает сеанс. ", "Как работает: удаляет токен на сервере и открывает окно входа.")]
DESC = ("Мой профиль", "Кто вы в команде (имя, логин, роль, привязка к участнику), ваша работа (открытые задачи по спринтам, баги, импедименты, риски, "
        "компетенции), ваше время (загрузка в спринте, отпуска, день рождения) и упоминания на доске. Здесь же тема, плотность, стартовая страница "
        "и — при работе с сервером — смена своего пароля; привязка к участнику хранится в учётной записи.")
API_ROWS = [("PUT /api/auth/me", "VIEWER+", "«Мой профиль»: привязка своей учётной записи к участнику команды (поле member; пустое — снять)."),
            ("PUT /api/auth/me/password", "VIEWER+", "Смена своего пароля: oldPassword, newPassword (не короче 8). 204 — успех, 403 — неверный текущий, 400 — короткий или прежний, 429 — много попыток; другие сеансы закрываются.")]


def put_before(anchor, p):
    anchor._p.addprevious(p._p); return p


def add_pic(d, anchor, img, cap, width):
    if not (SH / img).exists(): print("   ! нет снимка", img); return
    p = d.add_paragraph(); p.alignment = 1; p.add_run().add_picture(str(SH / img), width=Inches(width)); put_before(anchor, p)
    c = d.add_paragraph(cap); c.alignment = 1
    for r in c.runs: r.italic = True
    put_before(anchor, c)


def set_text(p, text):
    for r in p.runs[1:]: r.text = ""
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)


def replace_in(d, a, b):
    n = 0
    for p in d.paragraphs:
        if a in p.text and b not in p.text: set_text(p, p.text.replace(a, b)); n += 1
    return n


def version(d):
    pars = list(d.paragraphs) + [p for t in d.tables for r in t.rows for c in r.cells for p in c.paragraphs] \
        + [p for s in d.sections for hf in (s.header, s.footer) for p in hf.paragraphs]
    n = 0
    for p in pars:
        if OLDV in p.text:
            for r in p.runs: r.text = r.text.replace(OLDV, NEWV)
            n += 1
    return n


def heading(d, start, lvl="Heading 1"):
    return next((k for k, p in enumerate(d.paragraphs) if p.style.name == lvl and p.text.startswith(start)), None)


# ---------- руководство пользователя ----------
fn = f"Руководство_пользователя_{P}.docx"; d = docx.Document(fn)
if not any(p.text.startswith(CH) for p in d.paragraphs):
    anchor = d.paragraphs[heading(d, "Ширина столбцов")]
    put_before(anchor, d.add_paragraph(CH, style="Heading 1"))
    for title, text, bullets, pics in SECTIONS:
        put_before(anchor, d.add_paragraph(title, style="Heading 2")); put_before(anchor, d.add_paragraph(text))
        for b in bullets: put_before(anchor, d.add_paragraph(b, style="List Bullet"))
        for img, cap in pics: add_pic(d, anchor, img, cap, 6.3)
    ps = d.paragraphs; toc = next(k for k, p in enumerate(ps) if p.style.name == "Normal" and p.text.startswith("20. Навигация"))
    ps[toc + 1].insert_paragraph_before(CH, style=ps[toc].style)
    replace_in(d, UG_201_OLD, UG_201_NEW); replace_in(d, UG_CH2_OLD, UG_CH2_NEW)
    i = heading(d, "17.1 Верхняя панель", "Heading 2"); ps = d.paragraphs; k = i + 1; last = None
    while ps[k].style.name == "List Bullet":
        p = ps[k]
        for a, b in CTRL:
            if p.runs and p.runs[0].text == a: p.runs[0].text = b
        if p.runs and p.runs[0].text.startswith("«🏠 Стартовая страница»"): last = p
        k += 1
    last = last or ps[k - 1]
    for lab, what, how in reversed(CTRL_NEW):
        np_ = d.add_paragraph(style="List Bullet"); r = np_.add_run(lab); r.bold = True; np_.add_run(what); np_.add_run(how); last._p.addnext(np_._p)
    st = "добавлена"
else: st = "уже есть"
d.save(fn); print(fn[:58], "глава 21:", st)

# ---------- описания ----------
for fn, pics in ((f"Иллюстрированное_описание_{P}.docx", True), (f"Описание_{P}.docx", False)):
    d = docx.Document(fn)
    if not any(p.style.name.startswith("Heading") and p.text.strip() == DESC[0] for p in d.paragraphs):
        i = heading(d, "Мой день", "Heading 2")
        anchor = next(p for p in d.paragraphs[i + 1:] if p.style.name.startswith("Heading"))
        put_before(anchor, d.add_paragraph(DESC[0], style="Heading 2")); put_before(anchor, d.add_paragraph(DESC[1]))
        if pics: add_pic(d, anchor, "s_profile.jpg", "Мой профиль", 6.61)
        st = "добавлен"
    else: st = "уже есть"
    d.save(fn); print(fn[:58], "«Мой профиль»:", st)

# ---------- руководство администратора ----------
fn = f"Руководство_администратора_{P}.docx"; d = docx.Document(fn); n = 0
n += replace_in(d, "(миграции V1…V12)", "(миграции V1…V13)")
for p in d.paragraphs:
    if p.text.startswith("Основная БД planirovanie состоит из") and "V13" not in p.text:
        set_text(p, p.text.replace("Flyway V1…V12:", "Flyway V1…V13:").rstrip(".")
                 + ", V12 — диаграмма Исикавы (ishikawa, ishikawa_cats), V13 — привязка учётной записи к участнику команды (app_users.member_name)."); n += 1
    if p.text.startswith("v1.47 (миграция V12)") and not any(x.text.startswith("v1.51 (миграция V13)") for x in d.paragraphs):
        np_ = copy.deepcopy(p._p); p._p.addnext(np_)
        from docx.text.paragraph import Paragraph
        set_text(Paragraph(np_, p._parent), "v1.51 (миграция V13): колонка app_users.member_name — участник команды, с которым связана учётная запись («Мой профиль»)."); n += 1
    if p.text.startswith("Администратор управляет учётными записями") and "Свой пароль" not in p.text:
        p.add_run(" Свой пароль и привязку к участнику команды любой пользователь меняет сам в «Моём профиле» (PUT /api/auth/me/password, PUT /api/auth/me)."); n += 1
    if p.text.startswith("Категории событий:") and "password_changed" not in p.text:
        p.add_run(" События профиля: profile_updated (привязка к участнику) и password_changed (смена своего пароля) — категория AUTH."); n += 1
tb = next(t for t in d.tables if t.rows[0].cells[0].text == "Метод и путь" and any("/api/auth/me" in r.cells[0].text for r in t.rows))
if not any(r.cells[0].text == API_ROWS[0][0] for r in tb.rows):
    ref = next(r for r in tb.rows if r.cells[0].text == "GET /api/auth/me")
    for vals in reversed(API_ROWS):
        tr = copy.deepcopy(ref._tr); ref._tr.addnext(tr)
        from docx.table import _Row
        row = _Row(tr, tb)
        for c, v in zip(row.cells, vals): set_text(c.paragraphs[0], v)
    n += 1
d.save(fn); print(fn[:58], "правок:", n)

# ---------- инструкция по развёртыванию ----------
fn = f"Инструкция_по_развёртыванию_{P}.docx"; d = docx.Document(fn); n = 0
n += replace_in(d, "миграции V1–V7)", "миграции V1–V13)")
n += replace_in(d, "V7 (team.scrum — доп. роль «Скрам-мастер»).",
                "V7 (team.scrum — доп. роль «Скрам-мастер»), V8 (цели спринтов, WSJF), V9 (WIP, DoR, scope-лог, OKR), V10 (реестры управления проектом, иерархия работ), "
                "V11 (Team Health Check, деплои, инциденты), V12 (диаграмма Исикавы), V13 (app_users.member_name — «Мой профиль»).")
n += replace_in(d, "При старте Flyway применит миграции V1…V7.", "При старте Flyway применит миграции V1…V13.")
n += replace_in(d, "Не редактируйте уже применённые Flyway-миграции (V1…V7) — добавляйте только новые (V7, …).",
                "Не редактируйте уже применённые Flyway-миграции (V1…V13) — добавляйте только новые (V14, …).")
last = d.paragraphs[-1]
if last.text.startswith("v1.43:") and not any(p.text.startswith("v1.51:") for p in d.paragraphs):
    for txt in ("v1.47: при старте применяются миграции Flyway V1…V12 (V12 добавляет диаграмму Исикавы).",
                "v1.51: при старте применяются миграции Flyway V1…V13 (V13 добавляет app_users.member_name для «Моего профиля»)."):
        np_ = copy.deepcopy(d.paragraphs[-1]._p); d.paragraphs[-1]._p.addnext(np_); set_text(d.paragraphs[-1], txt)
    n += 1
d.save(fn); print(fn[:58], "правок:", n)

for fn in sorted(glob.glob(f"*_{P}.docx")):
    d = docx.Document(fn); n = version(d); d.save(fn); print(f"{fn[:58]:58s} версия: {n}")
