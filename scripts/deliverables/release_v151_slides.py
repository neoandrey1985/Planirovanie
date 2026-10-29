# -*- coding: utf-8 -*-
"""v1.51.0: слайд «Мой профиль» в Презентация_приложения.pptx после «Новый проект за один шаг». Оформление клонируется
со слайда «Дашборд», снимок — capture_special_shots.py (build/shots/s_profile.jpg). Идемпотентно.
Запуск: python scripts/deliverables/release_v151_slides.py"""
import copy, os, pathlib
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

os.chdir(pathlib.Path(__file__).resolve().parents[2])
DECK, SH = "Презентация_приложения.pptx", pathlib.Path("build/shots")
SLIDES = [
 ("👤", "Мой профиль", "Кто я, моя работа и время, настройки.", "s_profile.jpg", [
   "•  Привязка к участнику команды — с сервером на любом устройстве.",
   "•  Мои задачи, баги, импедименты, риски и компетенции.",
   "•  Загрузка в спринте, отпуска и упоминания на доске.",
   "•  Тема, плотность, стартовая страница и смена своего пароля."]),
]


def texts(s): return [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
def title_of(s):
    t = texts(s); return t[1].text_frame.text.strip() if len(t) > 1 else (t[0].text_frame.text.strip() if t else "")


def set_text(shape, text):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p'))
    for p in ps[1:]: tb.remove(p)
    runs = ps[0].findall(qn('a:r'))
    for r in runs[1:]: ps[0].remove(r)
    runs[0].find(qn('a:t')).text = text


def set_paras(shape, lines):
    tb = shape.text_frame._txBody; ps = tb.findall(qn('a:p')); tpl = copy.deepcopy(ps[0])
    for p in ps: tb.remove(p)
    for ln in lines:
        p = copy.deepcopy(tpl); runs = p.findall(qn('a:r'))
        for r in runs[1:]: p.remove(r)
        runs[0].find(qn('a:t')).text = ln; tb.append(p)


def parts(s):
    t = texts(s); icon = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame)
    body = next(sh for sh in t if sh.text_frame.text.strip().startswith("•"))
    rest = [sh for sh in t if sh._element is not icon._element and sh._element is not body._element]
    pic = next(sh for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return icon, rest[0], rest[1], body, pic


def clone_after(prs, src, pos, shot=None):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes): sh._element.getparent().remove(sh._element)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
    pic = None
    for sh in src.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: pic = sh; continue
        new.shapes._spTree.append(copy.deepcopy(sh._element))
    if pic is not None and shot: new.shapes.add_picture(str(SH / shot), pic.left, pic.top, pic.width, pic.height)
    lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(pos, el)
    return new


def set_picture(slide, pic, path):
    slide.part.related_part(pic._element.find('.//' + qn('a:blip')).get(qn('r:embed')))._blob = pathlib.Path(path).read_bytes()


prs = Presentation(DECK)
src = next(s for s in prs.slides if title_of(s) == "Дашборд")
pos0 = next(i for i, s in enumerate(prs.slides) if title_of(s) == "Новый проект за один шаг") + 1
for k, (ic, title, sub, shot, bullets) in enumerate(SLIDES):
    s = next((x for x in prs.slides if title_of(x) == title), None)
    if s is None: s = clone_after(prs, src, pos0 + k, shot)
    icon, tsh, ssh, bsh, psh = parts(s)
    set_text(icon, ic); set_text(tsh, title); set_text(ssh, sub); set_paras(bsh, bullets); set_picture(s, psh, SH / shot)
prs.save(DECK)
for i, s in enumerate(prs.slides, 1):
    if i <= 8: print(i, title_of(s))
print("слайдов:", len(prs.slides))
