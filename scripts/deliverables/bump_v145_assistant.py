# -*- coding: utf-8 -*-
"""v1.45.0 deliverable update: bump version/date strings in all .docx/.pptx and enrich the
assistant (чат-бот) description to reflect that it now knows every section and self-trains.
Idempotent: re-running only replaces the old version token if still present."""
import glob, os, docx
from pptx import Presentation

os.chdir(r"C:/Users/Андрей/Dropbox/Planirovanie")
OLDV, NEWV = "1.44.0", "1.45.0"
OLDD, NEWD = "05.10.2026", "06.10.2026"

ASSIST_ADD = (" Начиная с v1.45.0 помощник знает все разделы и функции приложения и отвечает на вопрос о любом из "
 "них («что такое Гант», «зачем Bus factor», «как работает DORA»), перечисляет разделы («какие разделы есть», "
 "«что в блоке Команда») и рассказывает, что нового. База знаний собирается автоматически из самого приложения, "
 "поэтому обновляется сама после каждого релиза.")
ASSIST_SHORT = (" С v1.45.0 знает все разделы и функции: отвечает на вопрос о любом разделе, перечисляет их и "
 "рассказывает про новинки; база знаний самообновляется после каждого релиза.")

def runwise_replace(doc):
    def fix(par):
        for r in par.runs:
            if OLDV in r.text: r.text = r.text.replace(OLDV, NEWV)
            if OLDD in r.text: r.text = r.text.replace(OLDD, NEWD)
    for p in doc.paragraphs: fix(p)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: fix(p)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            for p in hf.paragraphs: fix(p)

def append_sentence(par, text):
    if "знает все разделы" in par.text: return False  # already enriched
    r = par.add_run(text)
    return True

for fn in sorted(glob.glob("*.docx")):
    if fn.startswith("~$"): continue
    d = docx.Document(fn); runwise_replace(d)
    ps = d.paragraphs
    if "Руководство_пользователя" in fn:
        for p in ps:
            if p.text.startswith("Кнопка «💬»"): append_sentence(p, ASSIST_ADD); break
    if "Иллюстрированное_описание" in fn or fn.startswith("Описание_"):
        for p in ps:
            if p.text.startswith("Встроенный ассистент по данным проекта"): append_sentence(p, ASSIST_SHORT); break
    d.save(fn); print("docx:", fn)

for fn in sorted(glob.glob("*.pptx")):
    if fn.startswith("~$"): continue
    prs = Presentation(fn)
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            tf = sh.text_frame
            for para in tf.paragraphs:
                for r in para.runs:
                    if OLDV in r.text: r.text = r.text.replace(OLDV, NEWV)
                    if OLDD in r.text: r.text = r.text.replace(OLDD, NEWD)
            if tf.text.strip().startswith("•  Отвечает по живым данным"):
                tf.text = ("•  Знает все 56 разделов и функций: отвечает на вопрос о любом («что такое Гант», «зачем Bus factor»).\n"
                           "•  Перечисляет разделы и блоки, рассказывает, что нового; отвечает по живым данным (velocity, баги, бюджет).\n"
                           "•  Открывает нужный раздел и выполняет команды; работает офлайн. База знаний самообновляется после релиза.")
    prs.save(fn); print("pptx:", fn)
print("done v1.45.0 bump + assistant enrich")
