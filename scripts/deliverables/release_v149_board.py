# -*- coding: utf-8 -*-
"""v1.49.0 deliverables: доска, волна 2 (проведение сессий + схемы и диаграммы).
Поднимает версию, дописывает разделы о доске в документах и на слайде, вставляет второй снимок доски
(build/shots/board2.jpg — см. capture_board_v149.py). Идемпотентно: повторный запуск ничего не дублирует."""
import glob, os, copy, pathlib, docx
from docx.shared import Inches
from docx.oxml.ns import qn
from pptx import Presentation
os.chdir(pathlib.Path(__file__).resolve().parents[2])
OLDV, NEWV = "1.48.0", "1.49.0"
SHOT2 = "build/shots/board2.jpg"
MARK = "v1.49"
NOTE = ("Новое в v1.49: проведение сессий — голосование точками с лимитом голосов, скрытым режимом и раскрытием итогов "
        "(медали 🥇🥈🥉), анонимный ввод, показ по фреймам как по слайдам и режим «следовать за ведущим», комментарии к "
        "элементам с @упоминаниями; схемы — фигуры блок-схем (начало/конец, процесс, решение, ввод/вывод, документ, база "
        "данных), подписи стрелок, майнд-карта (Tab — ветка, Enter — соседняя), таблица на холсте, дорожки ответственности "
        "по ролям команды и авто-раскладка выделенного сеткой или деревом.")
CAPTION = "Доска (v1.49): блок-схема с подписями стрелок, майнд-карта, таблица и итоги голосования"
UG_58 = [
 "Сессия (🗳): голосование точками — ведущий задаёт число голосов на участника и включает скрытый режим; клик по 👍 на стикере отдаёт голос, Shift+клик снимает. «Раскрыть итоги» показывает суммы и медали 🥇🥈🥉 у трёх лидеров, «Завершить» фиксирует результат, «Сбросить» обнуляет голоса.",
 "Анонимный ввод: пока он включён, авторы стикеров скрыты (@🎭 аноним); «Раскрыть авторов» показывает их после обсуждения.",
 "Показ по фреймам: фреймы верхнего уровня становятся слайдами в порядке чтения (слева направо, сверху вниз); ← → или пробел — листать, ⛶ — во весь экран, Esc — выход. Участники с отметкой «следовать за ведущим» видят тот же слайд у себя.",
 "Комментарии (💬): выделите элемент и нажмите 💬 на панели действий; @имя упоминает участника команды (кнопки-подсказки под полем). На элементе появляется значок с числом комментариев, сводка всех обсуждений — в панели «Сессия».",
 "Схемы (🔷): фигуры блок-схем — начало/конец, процесс, решение, ввод/вывод, документ, база данных; двойной клик по стрелке — подпись. Майнд-карта: Tab — дочерняя ветка, Enter — соседняя, ветки раскладываются автоматически в обе стороны. Таблица: двойной клик по ячейке — правка (Tab/Enter — следующая ячейка), строки и столбцы — кнопками на панели действий. Дорожки ответственности строятся по ролям команды.",
 "Авто-раскладка: выделите элементы и нажмите ⊞ (сеткой) или 🌳 (деревом по стрелкам, для майнд-карты — выровнять ветки).",
]
UG_1715 = [
 "«🔷 Схемы» — меню фигур блок-схем, майнд-карты, таблицы и дорожек ответственности. Как работает: фигура выбирается как инструмент — кликните или потяните по холсту.",
 "«🗳 Сессия» — голосование точками, анонимный ввод, показ по фреймам, «следовать за ведущим» и сводка обсуждений. Как работает: настройки сессии синхронизируются с участниками в реальном времени.",
 "Панель действий (дополнительно): 💬 комментарии, ⊞ сетка, 🌳 дерево; для таблицы — «+ строка», «+ столбец», «− строка», «− столбец».",
]

def fixruns(p):
    for r in p.runs:
        if OLDV in r.text: r.text = r.text.replace(OLDV, NEWV)

def all_pars(d):
    for p in d.paragraphs: yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: yield p
    for s in d.sections:
        for hf in (s.header, s.footer):
            for p in hf.paragraphs: yield p

def next_heading_idx(ps, i):
    for j in range(i + 1, len(ps)):
        if ps[j].style.name.startswith("Heading"): return j
    return None

def add_picture_before(d, anchor):
    """Вставить абзац с картинкой и подписью перед anchor (или в конец документа)."""
    pic = d.add_paragraph(); pic.alignment = 1; pic.add_run().add_picture(SHOT2, width=Inches(6.2))
    cap = d.add_paragraph(CAPTION); cap.alignment = 1
    if cap.runs: cap.runs[0].italic = True
    if anchor is not None:
        anchor._p.addprevious(pic._p); anchor._p.addprevious(cap._p)
    return 1

for fn in sorted(glob.glob("*.docx")) + sorted(glob.glob("Тестовая документация/*.docx")):
    if "~$" in fn: continue
    d = docx.Document(fn)
    for p in all_pars(d): fixruns(p)
    ps = d.paragraphs; added = imgs = 0
    done = any(p.text.startswith("Сессия (🗳)") for p in ps) or any(MARK in p.text and "проведение сессий" in p.text for p in ps)
    if fn.startswith("Руководство_пользователя") and not done:
        for title, lines, pic in (("5.8 Доска", UG_58, True), ("17.15 Доска", UG_1715, False)):
            i = next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and p.text.startswith(title)), None)
            if i is None: continue
            j = next_heading_idx(ps, i)
            for ln in lines:
                if j is not None: ps[j].insert_paragraph_before(ln, style="List Bullet")
                else: d.add_paragraph(ln, style="List Bullet")
                added += 1
            ps = d.paragraphs
            if pic:
                j = next_heading_idx(ps, i); imgs += add_picture_before(d, ps[j] if j is not None else None); ps = d.paragraphs
    if (fn.startswith("Иллюстрированное_описание") or fn.startswith("Описание_")) and not done:
        i = next((k for k, p in enumerate(ps) if p.style.name.startswith("Heading") and p.text.startswith("Доска")), None)
        if i is not None:
            j = next_heading_idx(ps, i)
            ref = next((p for p in ps[i + 1:j] if p.text.startswith("Новое в v1.4")), ps[i + 1])
            np_ = ps[j].insert_paragraph_before(NOTE) if j is not None else d.add_paragraph(NOTE)
            np_.style = ref.style; added += 1
            if fn.startswith("Иллюстрированное_описание"):
                ps = d.paragraphs; j = next_heading_idx(ps, i); imgs += add_picture_before(d, ps[j] if j is not None else None)
    d.save(fn); print("docx %-60s +%d text, +%d images" % (fn[:60], added, imgs))

prs = Presentation("Презентация_приложения.pptx")
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs: fixruns(p)
    texts = [sh for sh in s.shapes if sh.has_text_frame]
    if not any("Доска — интерактивный бесконечный холст" in sh.text_frame.text for sh in texts): continue
    body = next(sh for sh in texts if sh.text_frame.text.strip().startswith("•"))
    if MARK not in body.text_frame.text:
        last = body.text_frame.paragraphs[-1]._p
        np_ = copy.deepcopy(last); last.addnext(np_)
        runs = np_.findall(qn('a:r'))
        for r in runs[1:]: np_.remove(r)
        runs[0].find(qn('a:t')).text = "•  Сессии и схемы (v1.49): голосование точками, показ по фреймам, комментарии; блок-схемы, майнд-карты, таблицы."
    print("deck: board slide updated")
prs.save("Презентация_приложения.pptx"); print("done")
