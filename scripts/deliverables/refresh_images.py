# -*- coding: utf-8 -*-
"""Заменить снимки экрана в .docx и Презентация_приложения.pptx свежими из build/shots.

Каждая картинка в документе опознаётся по SHA-1 её содержимого: shot_hashes.json хранит «хэш → имя снимка» для
всех снимков, когда-либо размещённых в документах (плюс --old <папка> со старыми снимками). Если для имени есть свежий build/shots/<имя>.jpg — картинка заменяется (одна часть может быть
общей для нескольких слайдов — тогда обновляются все). Неопознанные картинки (схемы, диаграммы) не трогаются.
Запуск: python scripts/deliverables/refresh_images.py [--old <папка со старыми снимками>]
"""
import argparse, glob, hashlib, json, os, pathlib, docx
from pptx import Presentation

os.chdir(pathlib.Path(__file__).resolve().parents[2])
MAP = pathlib.Path("scripts/deliverables/shot_hashes.json")
SH = pathlib.Path("build/shots")
ap = argparse.ArgumentParser(); ap.add_argument("--old", action="append", default=[]); args = ap.parse_args()

known = json.loads(MAP.read_text(encoding="utf-8")) if MAP.exists() else {}
for d in args.old:
    for p in pathlib.Path(d).glob("*.jpg"): known.setdefault(hashlib.sha1(p.read_bytes()).hexdigest(), p.stem)
fresh = {p.stem: p.read_bytes() for p in SH.glob("*.jpg")}
for name, data in fresh.items(): known[hashlib.sha1(data).hexdigest()] = name


def refresh(parts, label):
    done = same = unknown = 0
    for part in parts:
        name = known.get(hashlib.sha1(part.blob).hexdigest())
        if name is None or name not in fresh: unknown += 1; continue
        if part.blob == fresh[name]: same += 1; continue
        part._blob = fresh[name]; done += 1
    print(f"{label[:60]:60s} заменено {done:3d}, уже свежих {same:3d}, не опознано {unknown}")


for fn in sorted(glob.glob("*.docx")):
    if fn.startswith("~$"): continue
    d = docx.Document(fn)
    parts = {id(r.target_part): r.target_part for r in d.part.rels.values() if "image" in r.reltype}.values()
    refresh(parts, fn); d.save(fn)
prs = Presentation("Презентация_приложения.pptx")
parts = {}
for s in prs.slides:
    for r in s.part.rels.values():
        if "image" in r.reltype: parts[r.target_part.partname] = r.target_part
refresh(parts.values(), "Презентация_приложения.pptx"); prs.save("Презентация_приложения.pptx")
MAP.write_text(json.dumps(known, ensure_ascii=False, indent=0, sort_keys=True), encoding="utf-8")
